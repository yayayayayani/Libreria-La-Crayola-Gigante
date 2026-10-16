"""Genera los seis Word desde fuentes editables. Requiere python-docx.
En este equipo: python docs/entrega2/generar_documentos.py
"""
from pathlib import Path
import sys, re, json
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / '.tools/python'))
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

BASE = Path(__file__).resolve().parent
OUT = BASE / 'Word'
OUT.mkdir(exist_ok=True)

def clean(text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    return text.replace('**', '').replace('`', '')

def field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    code = OxmlElement('w:instrText'); code.set(qn('xml:space'), 'preserve'); code.text = instruction
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, code, end])

def base_document(title):
    d = Document()
    section = d.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = Inches(.5)
    normal = d.styles['Normal']
    normal.font.name = 'Times New Roman'; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(.5)
    for key, alignment in [('Heading 1', WD_ALIGN_PARAGRAPH.CENTER), ('Heading 2', WD_ALIGN_PARAGRAPH.LEFT), ('Heading 3', WD_ALIGN_PARAGRAPH.LEFT)]:
        style = d.styles[key]; style.font.name = 'Times New Roman'; style.font.size = Pt(12)
        style.font.bold = True; style.font.color.rgb = RGBColor(0,0,0)
        style.paragraph_format.alignment = alignment
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.space_before = Pt(12); style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 2; style.paragraph_format.keep_with_next = True
    for key in ['Title','Subtitle','Caption']:
        d.styles[key].font.name='Times New Roman'
        d.styles[key].font.color.rgb=RGBColor(0,0,0)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.first_line_indent = Inches(0)
    field(header, 'PAGE')
    for _ in range(3): d.add_paragraph('')
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0)
    p.add_run(title).bold=True
    d.add_paragraph('')
    for text in ['David [apellidos], Yordy [apellidos] y Daylin [apellidos]', '[Institución y carrera]', '[Curso y sección]', '[Nombre del docente]', '16 de octubre de 2026', 'Librería La Crayola Gigante · Entrega 2']:
        p=d.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0)
    d.add_page_break()
    d.add_heading('Contenido',1)
    field(d.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    p=d.add_paragraph('Índice automático: en Word, hacer clic derecho y elegir Actualizar campo si cambia la paginación.');p.paragraph_format.first_line_indent=Inches(0);p.runs[0].italic=True
    d.add_page_break()
    d.core_properties.title=title;d.core_properties.author='Equipo: David, Yordy y Daylin'
    d.core_properties.subject='Entrega 2: compras e inventario'
    d.core_properties.comments='Documento editable. Completar datos académicos y evidencias reales antes de entregar.'
    return d

def caption(d, kind, number, title):
    p=d.add_paragraph();p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.keep_with_next=True
    p.add_run(f'{kind} {number}').bold=True
    p=d.add_paragraph();p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.keep_with_next=True
    p.add_run(title).italic=True

def table(d, rows, number, title):
    caption(d,'Tabla',number,title)
    t=d.add_table(rows=0, cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    widths=[6.5/len(rows[0])]*len(rows[0])
    if len(rows[0]) >= 4 and len(rows[0][0])<12:
        widths[0]=.65; widths[1:]=[(6.5-.65)/(len(widths)-1)]*(len(widths)-1)
    for col,width in zip(t.columns,widths): col.width=Inches(width)
    for index,row in enumerate(rows):
        cells=t.add_row().cells
        props=t.rows[-1]._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');props.append(no_split)
        if index==0:
            repeat=OxmlElement('w:tblHeader');props.append(repeat)
        for cell,value,width in zip(cells,row,widths):
            cell.width=Inches(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Inches(0)
            p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
            r=p.add_run(clean(value));r.font.size=Pt(10);r.bold=index==0
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','bottom'] if index==0 else (['bottom'] if index==len(rows)-1 else []):
                b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'6');borders.append(b)
            cell._tc.get_or_add_tcPr().append(borders)
    d.add_paragraph('')

def placeholder(d,ident,title,instruction,number):
    caption(d,'Figura',number,f'{title} — evidencia {ident}')
    t=d.add_table(rows=1,cols=1);t.alignment=WD_TABLE_ALIGNMENT.CENTER
    cell=t.cell(0,0)
    p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Inches(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f'INSERTAR AQUÍ LA CAPTURA {ident}').bold=True
    p=cell.add_paragraph(instruction);p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.line_spacing=1.2
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F2F2');cell._tc.get_or_add_tcPr().append(shade)
    p=cell.add_paragraph('Fecha y hora: __________   Capturó: __________\nResultado observado: ____________________________________')
    p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.line_spacing=1.5
    p=d.add_paragraph('Nota. Espacio reservado para evidencia real del equipo; aún no incorporada.');p.paragraph_format.first_line_indent=Inches(0);p.runs[0].italic=True

def figure(d,filename,title,number):
    image=BASE/'figuras'/filename
    if not image.exists(): image=BASE/'evidencias'/filename
    if not image.exists(): raise FileNotFoundError(image)
    caption(d,'Figura',number,title)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0)
    p.paragraph_format.line_spacing=1
    # All prepared landscape figures fit inside the printable area.
    p.add_run().add_picture(str(image),width=Inches(6.3))
    note='Nota. Elaboración técnica a partir del código y modelo de datos del proyecto.'
    if 'demo' in filename: note='Nota. Captura de prueba de interfaz con datos de demostración y respuestas API simuladas; no acredita una reunión ni una compra real del equipo.'
    p=d.add_paragraph(note);p.paragraph_format.first_line_indent=Inches(0);p.runs[0].italic=True

def expand(text):
    legacy=ROOT/'docs/sprint2'
    for key,file in [('BACKLOG','BACKLOG.md'),('REGLAS','REGLAS.md'),('PRUEBAS','PRUEBAS.md'),('MODELO','MODELO_Y_DIAGRAMAS.md')]:
        value=(legacy/file).read_text(encoding='utf-8')
        if key=='MODELO': value=value.split('## Diccionario del incremento')[1].split('## Arquitectura compartida')[0]
        if key=='REGLAS': value=value.split('## NIT y CUI')[0].replace('# Reglas de negocio','')
        if key=='BACKLOG': value=value.split('Responsables:')[0]
        if key=='PRUEBAS': value=value.split('## Validaci')[0]
        value=re.sub(r'^# ', '### ',value,flags=re.M)
        text=text.replace('{{'+key+'}}',value)
    results=(BASE/'RESULTADOS.md').read_text(encoding='utf-8')
    return text.replace('{{RESULTADOS}}',results)

def render(source):
    text=expand(source.read_text(encoding='utf-8'))
    lines=text.splitlines();title=lines.pop(0).lstrip('# ')
    d=base_document(title);d.add_heading(title,1)
    i=0;nt=0;nf=0;section='Contenido';in_code=False;references=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line: continue
        if line.startswith('```'):
            in_code=not in_code;continue
        if in_code:
            p=d.add_paragraph(clean(line));p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.line_spacing=1.2
            for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)
            continue
        if line.startswith('[['):
            parts=line[2:-2].split('|')
            kind,ident=parts[0].split(':',1);nf+=1
            if kind=='CAPTURA':placeholder(d,ident,parts[1],parts[2],nf)
            else:figure(d,ident,parts[1],nf)
            continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'));section=line.lstrip('# ')
            references=section=='Referencias'
            d.add_heading(clean(section),min(max(level-1,1),3));continue
        if line.startswith('|'):
            block=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i].strip());i+=1
            rows=[[v.strip() for v in row.strip('|').split('|')] for row in block if not re.match(r'^\|[\s:|-]+\|$',row)]
            nt+=1;table(d,rows,nt,section);continue
        p=d.add_paragraph(clean(line))
        if references:
            p.paragraph_format.left_indent=Inches(.5);p.paragraph_format.first_line_indent=Inches(-.5)
        if line.startswith('- '):p.paragraph_format.first_line_indent=Inches(0)
    output=OUT/(source.stem+'.docx');d.save(output)
    # Reopen the zip with the same parser; malformed Word packages fail here.
    check=Document(output)
    assert len(check.paragraphs)>30 and check.sections[0].left_margin==Inches(1)
    return {'archivo':output.name,'parrafos':len(check.paragraphs),'tablas':len(check.tables),'figuras':len(check.inline_shapes),'bytes':output.stat().st_size}

if __name__=='__main__':
    report=[render(p) for p in sorted((BASE/'fuentes').glob('*.md'))]
    (BASE/'VALIDACION_WORD.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    for item in report: print(item['archivo'],item['bytes'],'bytes')
