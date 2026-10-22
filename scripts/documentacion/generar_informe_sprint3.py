"""Informe integrado Sprint 3. No sobrescribe el informe de referencia del Sprint 2."""
from pathlib import Path
import sys,json,re
from copy import deepcopy
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'.tools/docgen'))
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
BASE=Path(__file__).parent
OLD=json.loads((BASE/'base_documental.json').read_text(encoding='utf-8'))
d=Document(); sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
for name in ['Normal','Heading 1','Heading 2','Heading 3','Caption']:
 s=d.styles[name];s.font.name='Times New Roman';s.font.size=Pt(12);s.font.color.rgb=RGBColor(0,0,0)
 for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
  fonts=s.element.find('.//'+qn('w:rFonts'))
  if fonts is not None:fonts.attrib.pop(qn('w:'+attr),None)
 s.paragraph_format.line_spacing=2;s.paragraph_format.space_after=Pt(0)
 if name=='Normal':s.paragraph_format.first_line_indent=Inches(.5)
 if name.startswith('Heading'):
  s.font.bold=True;s.paragraph_format.first_line_indent=Inches(0);s.paragraph_format.keep_with_next=True
  s.paragraph_format.space_before=Pt(12)
  s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER if name=='Heading 1' else WD_ALIGN_PARAGRAPH.LEFT
def field(p,code):
 r=p.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin');r._r.append(a)
 t=OxmlElement('w:instrText');t.text=code;r._r.append(t)
 a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'end');r._r.append(a)
h=sec.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT;h.paragraph_format.first_line_indent=Inches(0);field(h,'PAGE')
def p(text):
 return d.add_paragraph(text)
def heading(text,level=1,page=False):
 if page:d.add_page_break()
 return d.add_heading(text,level)
nt=nf=0
def table(title,rows,widths=None):
 global nt
 nt+=1;caption('Tabla',nt,title)
 t=d.add_table(rows=0,cols=len(rows[0]));t.autofit=False
 widths=widths or [6.5/len(rows[0])]*len(rows[0])
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for i,row in enumerate(rows):
  cells=t.add_row().cells;pr=t.rows[-1]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for cell,value,w in zip(cells,row,widths):
   cell.width=Inches(w);z=cell.paragraphs[0];z.paragraph_format.first_line_indent=Inches(0)
   z.paragraph_format.line_spacing=1.12;z.paragraph_format.space_before=Pt(5);z.paragraph_format.space_after=Pt(5)
   r=z.add_run(str(value));r.font.name='Times New Roman';r.font.size=Pt(10);r.bold=i==0
   borders=OxmlElement('w:tcBorders')
   for edge in (['top','bottom'] if i==0 else ['bottom'] if i==len(rows)-1 else []):
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');borders.append(e)
   cell._tc.get_or_add_tcPr().append(borders)
 d.add_paragraph('')
 return t
def caption(kind,n,title):
 a=p(f'{kind} {n}');a.paragraph_format.first_line_indent=Inches(0);a.paragraph_format.keep_with_next=True;a.runs[0].bold=True
 a=p(title);a.paragraph_format.first_line_indent=Inches(0);a.paragraph_format.keep_with_next=True;a.runs[0].italic=True
def figure(file,title,note='Elaboración técnica a partir de la aplicación y su modelo de datos.'):
 global nf
 nf+=1;caption('Figura',nf,title)
 a=p('');a.paragraph_format.first_line_indent=Inches(0);a.paragraph_format.line_spacing=1
 a.add_run().add_picture(str(file),width=Inches(6.45))
 a=p('Nota. '+note);a.paragraph_format.first_line_indent=Inches(0);a.runs[0].italic=True
def slot(code,title,instruction):
 heading(title,2)
 t=d.add_table(rows=1,cols=1);c=t.cell(0,0)
 z=c.paragraphs[0];z.paragraph_format.first_line_indent=Inches(0);z.add_run('Evidencia '+code).bold=True
 for text in [instruction,'Fecha: __________  Responsable: __________','Identificador de operación: __________  Resultado observado: __________']:
  z=c.add_paragraph(text);z.paragraph_format.first_line_indent=Inches(0);z.paragraph_format.line_spacing=1.3
 sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F3F5');c._tc.get_or_add_tcPr().append(sh)
 p('')
def oldsection(doc,start,end=None):
 active=False
 for style,text in OLD[doc]['paragraphs']:
  if text==start:active=True;continue
  if active and end and text==end:break
  if not active:continue
  if style.startswith('Heading'):heading(text,2 if style=='Heading 1' else 3)
  elif not re.match(r'^(Tabla|Figura) \d+$',text) and not text.startswith(('Nota.','Índice automático')):p(text)
def oldtable(doc,index,title,widths=None):table(title,OLD[doc]['tables'][index],widths)

# Portada institucional: datos tomados del documento proporcionado por el equipo.
ref=Document(next((ROOT/'docs').glob('Sprint_2*.docx')))
for text in [x.text for x in ref.paragraphs[:5]]:
 a=p(text);a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.first_line_indent=Inches(0);a.paragraph_format.line_spacing=1.25
logo=ref.paragraphs[4]._p.xpath('.//a:blip')
if logo:
 part=ref.part.related_parts[logo[0].get(qn('r:embed'))];f=BASE/'figuras/logo.png';f.write_bytes(part.blob)
 a=p('');a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.first_line_indent=Inches(0);a.add_run().add_picture(str(f),width=Inches(1.25))
a=p('Sprint 3: Pedidos, ventas y procesos avanzados');a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.first_line_indent=Inches(0);a.runs[0].bold=True
a=p('Librería La Crayola Gigante');a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.first_line_indent=Inches(0)
for name in ['David Giovanni García Yoc','Yordy Misael Carreto López','Daylin Yanabell Loarca Escobar']:
 a=p(name+' | Carné: ______________');a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.first_line_indent=Inches(0);a.paragraph_format.line_spacing=1.5
a=p('21 de octubre de 2026');a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.first_line_indent=Inches(0)
d.add_page_break();a=p('Contenido');a.runs[0].bold=True;a.alignment=WD_ALIGN_PARAGRAPH.CENTER
field(p(''),'TOC \\o "1-2" \\h \\z \\u')

heading('1. Descripción general del incremento',page=True)
p('El tercer Sprint incorpora la atención de pedidos y el cobro de ventas a la operación de Librería La Crayola Gigante. El incremento parte del catálogo, los proveedores y las entradas de inventario del Sprint 2 y completa el recorrido comercial: identificar al cliente, registrar su pedido, comprobar existencias, resolver la aprobación cuando corresponda, generar la venta, recibir el pago y descontar las unidades entregadas.')
p('El problema que se atiende es la separación entre vender, cobrar y saber cuánto dinero tiene el negocio. Una venta registrada sin evidencia de pago no representa necesariamente un ingreso; del mismo modo, un pedido no debe consumir inventario antes de que se complete el cobro. El diseño distingue estos eventos y los relaciona mediante identificadores, detalles y movimientos consultables.')
p('La solución se orienta a una caja atendida por personal de la librería. El operador registra cómo pagó el cliente y el sistema calcula importes, IVA incluido, descuentos y puntos. La aplicación no procesa tarjetas bancarias ni confirma transferencias por sí misma: conserva el registro del cobro que el operador declara haber recibido.')
heading('1.1 Objetivo general',2)
p('Implementar un proceso de venta trazable que conecte pedidos, aprobaciones, cobros e inventario, con controles de acceso y consultas que permitan al administrador revisar la operación y sus movimientos financieros.')
heading('1.2 Objetivos específicos',2)
for text in ['Registrar pedidos con productos activos, cantidades válidas y precios determinados por el servidor.','Controlar los estados del pedido y separar la autorización de descuentos de la actuación del vendedor.','Registrar pagos completos, abonos y combinaciones de medios de pago sin duplicar operaciones.','Actualizar existencias y puntos al cerrar el cobro, preservando la integridad de los registros relacionados.','Proporcionar comprobantes internos, reportes descargables, auditoría y un panel financiero restringido a administración.','Verificar el comportamiento mediante pruebas automatizadas y una demostración funcional reproducible.']:p(text)
heading('1.3 Alcance y decisiones',2)
table('Alcance del Sprint 3',[['Área','Resultado incorporado','Límite del incremento'],['Venta','Pedido, detalle, aprobación, venta y cobro','El pedido no reserva inventario'],['Caja','Efectivo, tarjeta, transferencia y puntos','La confirmación bancaria se realiza fuera del sistema'],['Cliente','Registro desde caja, consumidor final y NIT sin cuenta','Solo clientes registrados acumulan puntos'],['Comprobante','Consulta, desglose de IVA e impresión','Documento interno; certificación FEL pendiente'],['Finanzas','Ingresos, egresos, saldo y pagos a proveedores','No sustituye conciliación bancaria ni contabilidad completa'],['Personal','Roles, carnet, sueldo, horario y ventas','No constituye una liquidación de nómina']],[1,2.7,2.8])
heading('1.4 Continuidad con los Sprints anteriores',2)
p('Se mantienen el inicio de sesión, la administración de usuarios, los permisos, el catálogo y las compras recibidas. El Sprint 3 utiliza esas relaciones: la entrada por compra aumenta existencias y la salida por venta pagada las disminuye. La incorporación de un rol Dueño amplía la organización del acceso administrativo sin convertir las cuentas de clientes en usuarios del sistema.')

heading('2. Organización del equipo y planificación',page=True)
p('La distribución de responsabilidades permite revisar el incremento en tres bloques complementarios. Cada responsable prepara su demostración y contrasta el resultado con otro integrante. La asignación expresa responsabilidad de revisión y entrega; la aceptación se registra al presentar las evidencias.')
oldtable('01',0,'Responsabilidades de David, Yordy y Daylin',[1,1.9,1.8,1.8])
heading('2.1 Product Backlog actualizado',2)
p('Se conservan las capacidades de acceso, catálogo y compras como base del producto. La prioridad del Sprint 3 es cerrar el proceso de venta. Las ampliaciones de caja, puntos y finanzas responden a las necesidades de operación planteadas durante la revisión del sistema.')
oldtable('01',1,'Product Backlog del producto',[.65,3.25,.7,1.9])
heading('2.2 Historias de usuario',2)
oldsection('01','Historias de usuario y aceptación','Sprint Backlog asignado')
heading('2.3 Sprint Backlog',2)
p('El objetivo de trabajo es que una operación pueda seguirse desde el cliente hasta el pago y su salida de inventario. Los puntos indican complejidad relativa para organizar la revisión; no representan horas efectivamente trabajadas.')
oldtable('01',2,'Tareas del Sprint y responsables',[1.55,1.45,2.85,.65])
heading('2.4 Tablero Kanban de cierre',2)
table('Tablero de seguimiento del incremento',[['Implementado','En revisión de entrega','Pendiente de aceptación'],['Pedidos y estados; aprobación; ventas y detalle','Recorrido completo con datos relacionados','Capturas del equipo y observaciones del docente'],['Pagos, puntos, IVA y comprobantes','Revisión de caja, comprobante e impresión','Confirmación del formato de impresión utilizado'],['Inventario, auditoría y permisos','Conciliación de saldos y revisión por rol','Evidencia de acceso desde los equipos del grupo'],['Reportes CSV y panel financiero','Comprobación de filtros y archivos descargados','Decisiones de Sprint Review y retrospectiva']])
heading('2.5 Criterio de terminado',2)
p('Una función se considera terminada técnicamente cuando valida las entradas, respeta permisos, conserva relaciones, controla los errores esperados y cuenta con verificación del comportamiento. El cierre de entrega añade la demostración, la documentación y la aceptación del equipo. Una captura aislada no sustituye la relación entre pedido, venta, pago y movimiento de inventario.')

heading('3. Diseño de la solución',page=True)
heading('3.1 Arquitectura',2)
p('React presenta la caja, los formularios y los listados. La API Express autentica al usuario mediante JWT, consulta sus permisos y aplica las reglas del negocio. PostgreSQL en Neon conserva los registros compartidos y las restricciones de integridad. El navegador ayuda a prevenir errores, pero los importes y autorizaciones se determinan en el servidor.')
figure(BASE/'figuras/image1.png','Arquitectura de la aplicación')
heading('3.2 Flujo de pedido, venta y pago',2)
p('El vendedor selecciona cliente o consumidor final y agrega los productos. El servidor consulta sus precios, comprueba existencias y calcula el total. Cuando el descuento supera el umbral o el importe exige revisión, el pedido queda pendiente. Una cuenta autorizada distinta del creador decide con un motivo. El pedido aprobado puede convertirse en venta; los abonos reducen el pendiente sin producir salida de inventario.')
figure(BASE/'figuras/image2.png','Proceso comercial del Sprint 3')
p('Al completar el pago se vuelve a comprobar el inventario. La operación registra el cobro, descuenta unidades, actualiza puntos cuando corresponde, guarda la auditoría y genera el comprobante. Si alguna condición impide finalizar, se revierte la transacción en lugar de conservar una venta pagada sin sus movimientos.')
heading('3.3 Estados y transiciones',2)
figure(BASE/'figuras/image3.png','Estados del pedido y decisiones permitidas')
oldtable('02',0,'Significado y acciones de cada estado',[1.65,2.35,2.5])
heading('3.4 Modelo relacional',2)
p('El pedido conserva su creador y puede relacionarse con un cliente registrado. Cuando se atiende a consumidor final o a un receptor con NIT sin cuenta, la ausencia de cliente es una condición prevista del modelo. No se trata de un detalle huérfano: los productos, pagos, aprobaciones y movimientos mantienen sus referencias a las operaciones que los originan.')
figure(BASE/'figuras/image4.png','Relaciones principales del proceso comercial')
oldtable('02',9,'Tablas de apoyo y relación con el flujo',[1.7,4.8])
heading('3.5 Integridad, concurrencia y trazabilidad',2)
oldsection('02','Transacciones e integridad','Referencias')
p('Las relaciones de auditoría que identifican una tabla y un registro son referencias de seguimiento, no claves foráneas universales. Por ello, la conservación del historial y las reglas que impiden borrar usuarios vinculados son parte del control del sistema. No se permite eliminar la cuenta propia ni retirar operaciones comerciales mediante el mantenimiento de usuarios.')
heading('3.6 Contratos de API y manejo de errores',2)
oldtable('02',12,'Operaciones de la API y controles de acceso',[2.2,2.25,2.05])
p('Los errores de validación explican qué dato debe corregirse. La falta de autenticación produce una respuesta 401; la ausencia de autorización produce 403; un conflicto de estado o existencias se comunica sin confirmar el cobro. Ante un reintento se utiliza la clave de operación para evitar duplicados. Un error de conexión no debe interpretarse como una confirmación de guardado.')

heading('4. Reglas de negocio y operación de caja',page=True)
oldtable('02',10,'Reglas verificables del Sprint 3',[.65,3.25,2.6])
heading('4.1 Precios, descuentos e IVA incluido',2)
p('Los precios registrados representan el importe final. El IVA no se agrega nuevamente al cobrar: se separa del total con impuesto incluido. Para una compra de Q112 sin descuento, la base imponible es Q100 y el IVA es Q12. Con un descuento de 10 %, el cliente paga Q100.80, distribuidos en Q90 de base y Q10.80 de IVA. Los cálculos monetarios del servidor evitan que la interfaz imponga un precio diferente.')
heading('4.2 Programa de puntos',2)
p('Por cada Q100 de pago monetario elegible se acumulan 10 puntos. El canje se realiza en bloques de 100 puntos, equivalentes a Q10. Un cliente con 100 puntos y una compra de Q110 puede utilizar Q10 en puntos y pagar Q100 en efectivo, tarjeta o transferencia. Los puntos canjeados no se registran como ingreso de dinero; su movimiento se conserva por separado para explicar el saldo del cliente.')
p('La identificación tributaria y la afiliación al programa tienen finalidades diferentes. Una persona puede proporcionar NIT para su comprobante sin autorizar la creación de una cuenta comercial. En ese caso, el receptor se guarda en la operación y no acumula puntos. Los datos del cliente registrado permiten consultar su historial y saldo sin exigirle una cuenta de acceso a la aplicación.')
heading('4.3 Cobros y comprobantes',2)
p('La caja admite efectivo, tarjeta y transferencia; solicita referencia para los métodos que la requieren. Si el cliente paga en dos momentos, cada abono queda visible. Solo cuando se cubre el total se completa la salida de inventario y se habilita el comprobante. El documento conserva los detalles y pagos de la venta para que modificaciones posteriores del catálogo no alteren el registro emitido.')
p('La impresión utiliza el navegador. El operador puede preparar la copia del cliente y la del negocio, o guardar el documento como PDF mediante el diálogo de impresión. La aplicación permite consultar los comprobantes dentro del historial. El documento es interno: mostrar IVA no lo convierte en un documento tributario certificado por FEL.')
heading('4.4 Finanzas: ingresos, egresos y saldo',2)
p('Mi dinero reúne los cobros y los movimientos financieros manuales. Permite consultar día, mes, año o un intervalo. El balance del período compara sus ingresos y egresos; el saldo disponible incorpora el historial anterior hasta la fecha final. Los cobros con puntos se excluyen del dinero recibido y los saldos de tarjeta y transferencia representan registros operativos, no conciliaciones bancarias.')
p('Una compra pagada a un proveedor genera una salida vinculada. Los insumos que no pertenecen al inventario se registran como gasto del negocio. Si el administrador vuelve a introducir como gasto un pago de compra ya registrado, duplicaría el egreso; el manual distingue estos recorridos. El saldo inicial se utiliza una vez por método y corresponde al inicio del historial.')
heading('4.5 Restricciones por rol y datos del empleado',2)
oldtable('02',11,'Funciones visibles y autorizadas por rol',[2.7,1.6,2.2])
p('La vista del empleado incluye su identificación, puesto, sueldo, horario y ventas realizadas. Los reportes financieros globales, las decisiones de aprobación y la administración del personal permanecen restringidos. Administrador y Dueño comparten las funciones administrativas previstas en esta versión; no se afirma que exista una jerarquía adicional de permisos entre ambos.')

heading('5. Verificación y casos de prueba',page=True)
p('La verificación combina pruebas de integración de API con PostgreSQL y pruebas de interfaz en navegador. Los resultados siguientes corresponden a la revisión técnica del 20 de octubre de 2026. Las pruebas de navegador utilizaron respuestas de API simuladas; por tanto, se distinguen de las capturas de la aplicación conectada incluidas en este informe.')
oldtable('03',0,'Resultados de la revisión automatizada',[1.5,2.2,2.8])
heading('5.1 Matriz de pruebas automatizadas',2)
oldtable('03',1,'Casos del backend y resultado registrado',[.55,4.85,1.1])
heading('5.2 Casos funcionales de demostración',2)
table('Guion de aceptación funcional',[['Caso','Preparación y acción','Resultado esperado'],['CP01. Venta disponible','Seleccionar producto con 10 unidades y vender 2; completar pago','Pedido pagado, comprobante y stock final 8'],['CP02. Stock insuficiente','Solicitar más unidades de las disponibles','Rechazo sin nuevo pago ni salida'],['CP03. Aprobación','Crear pedido con descuento de 20 %; decidir desde otro usuario','Bloqueo previo y autorización con motivo'],['CP04. Descuento','Aplicar 10 % sobre Q112','Total Q100.80, base Q90 e IVA Q10.80'],['CP05. Abonos','Pagar Q5 y después Q15 sobre una venta de Q20','Stock intacto tras primer abono; salida al completar'],['CP06. Pago excesivo','Enviar un importe superior al pendiente','Rechazo y saldo sin modificación'],['CP07. Pago combinado','Dividir una venta entre efectivo y tarjeta','Dos registros cuya suma cubre el total'],['CP08. Puntos','Canjear 100 puntos y cubrir la diferencia con dinero','Q10 de canje, reducción de puntos y dinero separado'],['CP09. NIT sin cuenta','Capturar receptor sin registrar cliente','Comprobante identificado sin alta ni puntos'],['CP10. Restricción de rol','Intentar acceder a finanzas como empleado','Función oculta y API denegada'],['CP11. Auditoría','Consultar las acciones del pedido pagado','Creador, decisión, pago y salida relacionados'],['CP12. Reportes','Consultar ventas y productos; descargar CSV','Archivos con encabezados y datos del filtro'],['CP13. Egreso','Registrar gasto de Q135 en efectivo','Egreso visible y saldo reducido en Q135'],['CP14. Formulario pendiente','Editar un dato e intentar cambiar de módulo','Guardar, descartar o continuar editando']],[1.25,2.75,2.5])
heading('5.3 Ejecución y alcance de las pruebas',2)
for text in ['Backend: npm run test:sprint3, desde la carpeta backend.','Frontend: npx playwright test tests/sprint3.spec.js tests/finanzas.spec.js tests/operacion.spec.js tests/permisos.spec.js.','Calidad de frontend: npm run lint y npm run build.','Las pruebas de integración crean registros temporales. Deben ejecutarse en una base destinada a pruebas cuando se comience a registrar información real del negocio.']:p(text)
heading('5.4 Datos de demostración y conciliación',2)
p('La aplicación dispone de un catálogo de demostración con artículos escolares, productos de oficina y materiales artísticos. Las compras abastecen el catálogo y las ventas utilizan esas existencias. Los clientes y empleados del conjunto son ficticios; los comprobantes y movimientos sirven para revisar el funcionamiento, no para representar operaciones reales del negocio.')
report=BASE/'resultado_carga.json'
if report.exists():
 data=json.loads(report.read_text(encoding='utf-8-sig'))
 table('Volumen cargado y comprobaciones de integridad',[['Indicador','Resultado']]+[[k,str(v)] for k,v in data.items()],[3.8,2.7])
p('La conciliación compara el stock de cada producto con sus entradas menos sus salidas. También se contrastan pagos y total de ventas pagadas, la suma de movimientos de puntos y el saldo del cliente. El conjunto incluye pedidos pendientes, una cancelación, una decisión de rechazo y una venta con abono para mostrar estados distintos del cierre exitoso.')
table('Lectura financiera del conjunto de demostración',[['Concepto','Importe','Interpretación'],['Saldo inicial','Q40,000.00','Q15,000 en efectivo y Q25,000 en transferencia'],['Ingresos monetarios','Q7,038.55','Cobros registrados; los puntos no son dinero'],['Egresos','Q30,930.00','Pagos de compras y gastos del negocio'],['Balance del período','−Q23,891.45','Ingresos menos egresos del período'],['Saldo disponible','Q16,108.55','Saldo inicial más ingresos menos egresos']],[1.5,1.25,3.75])
p('El balance negativo del período no significa que el saldo disponible sea negativo: la carga incluye la compra inicial de existencias y una apertura de dinero. La distinción puede observarse en el dashboard y permite explicar por qué las tarjetas de balance y saldo muestran importes diferentes.')

heading('6. Evidencias funcionales',page=True)
p('Las capturas de esta sección corresponden a la aplicación ejecutada con datos de demostración. Para la presentación, conservar el mismo identificador al mostrar un pedido, sus pagos y la auditoría. Las imágenes de interfaz acreditan el estado mostrado; no acreditan por sí mismas una reunión del equipo.')
screens=[('dashboard','Vista administrativa y resumen financiero'),('productos','Catálogo de productos de la librería'),('clientes','Clientes y programa de puntos'),('pedidos','Pedidos registrados y sus estados'),('reportes','Reporte inicial de ventas'),('comprobantes','Consulta de comprobantes internos'),('inventario','Compras y existencias disponibles'),('empleado','Vista del empleado y carnet')]
for file,title in screens:
 f=BASE/'capturas'/(file+'.png')
 if f.exists():figure(f,title,'Captura de la aplicación conectada a PostgreSQL, con datos de demostración, 21 de octubre de 2026.')
for code,title,instruction in [('E01','Venta completa e inventario','Colocar la captura del pedido pagado y el mismo producto antes y después de la salida. Mostrar el identificador de venta.'),('E02','Intento sin inventario','Solicitar una cantidad superior al stock y capturar el mensaje de rechazo; mostrar que no se registró pago.'),('E03','Decisión de aprobación','Mostrar el pedido pendiente y su decisión desde una cuenta diferente, con motivo visible.'),('E04','Descuento, IVA y pago','Mostrar el detalle de un cobro con subtotal, descuento, base imponible, IVA y método de pago.'),('E05','Auditoría del flujo','Filtrar la bitácora por el pedido utilizado en E01 y mostrar pago y salida de inventario.'),('E06','Descarga de reportes','Abrir los CSV de ventas y productos en una hoja de cálculo, conservando encabezados y rango consultado.')]:slot(code,title,instruction)

heading('7. Seguimiento Scrum y cierre del Sprint',page=True)
p('El seguimiento se organiza alrededor del objetivo de cerrar una venta verificable. La distribución por temas permite revisar dependencias: la aprobación depende del pedido; el cobro depende de una venta autorizada; el inventario depende del pago completo. El siguiente registro resume el avance técnico disponible y orienta la revisión del equipo.')
heading('7.1 Registro de seguimiento técnico',2)
table('Seguimiento de tareas y dependencias',[['Bloque','Avance disponible','Responsable de revisión','Comprobación siguiente'],['Pedido y aprobación','Estados, cálculo y decisión independiente','David','Demostrar descuento y rechazo por falta de stock'],['Venta y cobro','Detalle, abonos y pago final','Yordy','Conciliar pagos, salida y duplicados'],['Caja y fidelización','NIT sin cuenta, puntos e IVA','David y Daylin','Mostrar compra combinada y comprobante'],['Finanzas y permisos','Saldo, egresos, roles y vista del empleado','Yordy','Comparar acceso administrativo y empleado'],['Presentación y entrega','Reportes CSV, manual y matriz de requisitos','Daylin','Completar evidencias y aceptación']],[1.15,1.9,1.35,2.1])
heading('7.2 Daily Scrum: registro de la sesión',2)
p('No se dispone de evidencia de Daily históricos para este Sprint. El registro de tareas anterior describe el producto revisado y no se presenta como acta de reuniones pasadas. El siguiente formato permite incorporar una sesión efectivamente realizada y su captura sin atribuirle avances o asistentes no confirmados.')
table('Acta de seguimiento',[['Dato','Registro'],['Fecha y hora','____________________________'],['Participantes','____________________________'],['Objetivo inmediato','Revisar el flujo de venta y resolver pendientes para la entrega'],['Avance desde el seguimiento anterior','____________________________'],['Impedimento y acción','____________________________'],['Responsable y próxima comprobación','____________________________']],[2,4.5])
slot('S01','Evidencia de seguimiento del equipo','Insertar una captura de la reunión realizada con los participantes y el tablero o documento compartido. Anotar la fecha de esa sesión.')
heading('7.3 Sprint Review',2)
p('La Review se centra en el incremento funcionando. David presenta la creación del pedido y la aprobación; Yordy muestra pagos, inventario y restricciones; Daylin presenta reportes, comprobantes y la navegación. La decisión de aceptación debe vincularse con el criterio de cada historia, no únicamente con la apariencia de las pantallas.')
table('Guion de Sprint Review',[['Demostración','Criterio de aceptación','Decisión y observación'],['Venta con existencias','Total correcto, pago completo y salida única','________________'],['Pedido sujeto a aprobación','Decisión de otro usuario y motivo registrado','________________'],['Puntos y métodos combinados','Canje correcto sin inflar ingresos monetarios','________________'],['Comprobante e IVA','Desglose y consulta posterior conservados','________________'],['Reportes descargables','CSV de ventas y productos según filtros','________________'],['Acceso del empleado','Sin finanzas globales ni aprobación administrativa','________________']],[1.9,2.9,1.7])
slot('S02','Evidencia de Sprint Review','Insertar la captura de la presentación del incremento. Registrar quién revisó, historias aceptadas y observaciones pendientes.')
heading('7.4 Sprint Retrospective',2)
p('El análisis del incremento identifica tres aprendizajes: separar pedido y pago evita confundir ventas pendientes con ingresos; conservar relaciones simplifica explicar un movimiento de inventario; y una interfaz más clara reduce retrocesos al registrar clientes y cobrar. La revisión también muestra que la documentación debe consolidarse en un informe coherente y que las evidencias de reuniones necesitan registrarse en el momento de realizarlas.')
table('Mejoras para discutir en la retrospectiva',[['Observación','Consecuencia','Acción propuesta','Responsable'],['Cambios de alcance en caja y finanzas','Mayor dependencia entre pantallas y reglas','Mantener una lista priorizada antes de incorporar funciones','David'],['Necesidad de repetir recorridos comerciales','Riesgo de validar pantallas aisladas','Usar una venta de referencia de extremo a extremo','Yordy'],['Documentación inicialmente fragmentada','Dificulta lectura y revisión de entrega','Conservar un informe integrado con numeración continua','Daylin'],['Falta de evidencia de seguimiento','No permite acreditar Daily históricos','Registrar la sesión y sus acuerdos el mismo día','Equipo']],[1.6,1.6,2.4,.9])
table('Acuerdos de retrospectiva',[['Acuerdo confirmado','Responsable','Fecha compromiso','Comprobación'],['________________','________________','________________','________________'],['________________','________________','________________','________________']])
slot('S03','Evidencia de retrospectiva','Insertar captura del tablero de acuerdos revisado por el equipo. Completar responsables y fechas después de acordarlos.')

heading('8. Manual técnico y de usuario',page=True)
heading('8.1 Entorno y conexión',2)
oldsection('05','Conexión y puesta en marcha','Preparación de un catálogo desde cero')
p('La carga de demostración se ejecuta de forma explícita mediante backend/scripts/cargarDemostracion.js --cargar, desde backend. Conserva las cuentas existentes, guarda un respaldo local y utiliza claves de operación para reanudar la carga. No debe ejecutarse como parte del inicio habitual ni confundirse con una importación de ventas reales.')
heading('8.2 Inicio y navegación',2)
p('Ingresar con una cuenta autorizada y revisar el menú. Administrador y Dueño pueden consultar el panel financiero y mantener el personal. Empleado dispone de las funciones permitidas y de su carnet. Si una función no está disponible, revisar el permiso asignado; modificar la dirección o la interfaz no autoriza una operación en la API.')
heading('8.3 Atención de una venta',2)
oldsection('05','Atención de una venta','Consultar comprobantes e imprimir')
heading('8.4 Consulta e impresión de comprobantes',2)
oldsection('05','Consultar comprobantes e imprimir','Reportes y búsquedas')
heading('8.5 Búsquedas y reportes',2)
oldsection('05','Reportes y búsquedas','Consultar ingresos, egresos y saldo')
heading('8.6 Registro de gastos y consulta del saldo',2)
oldsection('05','Consultar ingresos, egresos y saldo','Administración y eliminación')
heading('8.7 Administración de personal y conservación de historial',2)
oldsection('05','Administración y eliminación','Salir con datos pendientes')
heading('8.8 Formularios con cambios pendientes',2)
oldsection('05','Salir con datos pendientes','Solución de problemas frecuentes')
heading('8.9 Solución de problemas',2)
oldtable('05',3,'Problemas frecuentes y comprobaciones',[2.15,4.35])

heading('9. Dictamen técnico y cumplimiento',page=True)
p('El incremento cubre el flujo comercial obligatorio y dispone de dos reportes iniciales descargables. La verificación automatizada registrada supera los diez casos mínimos. La carga de demostración hace visible la relación entre catálogo, compra, pedido, cobro e inventario y permite revisar la interfaz con operaciones completas y pendientes.')
rows=deepcopy(OLD['06']['tables'][0])
for row in rows:
 for i,txt in enumerate(row):
  for a,b in [('Documento 01','Sección 2'),('Documento 02','Secciones 3 y 4'),('Documento 04','Sección 7')]:txt=txt.replace(a,b)
  row[i]=txt
table('Matriz de cumplimiento del Sprint 3',rows,[2,2.8,1.7])
heading('9.1 Evaluación para una librería real',2)
p('La aplicación ofrece una base funcional para controlar operaciones de una librería pequeña, pero el paso de demostración a operación real exige distinguir control administrativo y obligaciones fiscales. La primera prioridad es integrar la emisión fiscal requerida y definir la conciliación de caja. También debe planificarse cómo registrar devoluciones sin borrar ventas, cómo comprobar respaldos y cómo supervisar accesos y saldos.')
oldtable('06',3,'Mejoras priorizadas para la siguiente etapa',[1.3,2.35,2.85])
heading('9.2 Conclusiones',2)
p('El Sprint 3 transforma el registro de ventas en un proceso relacionado. El pedido define lo solicitado; la aprobación deja constancia de la excepción; el pago determina cuándo se cierra la operación; y el movimiento explica el cambio de inventario. Esta separación permite consultar estados sin confundir dinero recibido con ventas aún pendientes.')
p('La experiencia de caja incorpora las necesidades principales del negocio: registrar al cliente sin abandonar la venta, capturar un NIT sin crear cuenta, combinar medios de pago y consultar puntos. Los reportes y el panel financiero complementan la operación al permitir revisar ingresos, egresos y productos vendidos con filtros y archivos descargables.')
p('La entrega técnica queda acompañada por un informe único, pruebas, diagramas y un manual. La aceptación académica debe completarse con las evidencias del equipo y las decisiones de revisión. La siguiente etapa debe priorizar facturación fiscal, devoluciones, arqueo y recuperación de respaldos antes de ampliar funciones secundarias.')

heading('Referencias',page=True)
for text in ['American Psychological Association. (2024). Paper format. APA Style. https://apastyle.apa.org/style-grammar-guidelines/paper-format','Schwaber, K., & Sutherland, J. (2020). The Scrum Guide. https://scrumguides.org/scrum-guide.html','PostgreSQL Global Development Group. (s. f.). Transactions. PostgreSQL documentation. https://www.postgresql.org/docs/current/tutorial-transactions.html','Superintendencia de Administración Tributaria. (s. f.). Factura electrónica en línea. https://en.portal.sat.gob.gt/portal/efactura/']:
 a=p(text);a.paragraph_format.left_indent=Inches(.5);a.paragraph_format.first_line_indent=Inches(-.5)
heading('Anexo A. Diccionario del núcleo de datos',page=True)
p('El diccionario recoge los campos del esquema consultado para el Sprint 3. Las tablas de catálogo, clientes y personal conservan las relaciones de las entregas anteriores. Los importes generados de base imponible e IVA se obtienen del total y de la tasa almacenada.')
for i,name in enumerate(['pedidos','detalle_pedidos','aprobaciones','ventas','detalle_ventas','pagos','movimientos_inventario','bitacora'],1):
 heading(name,2);oldtable('02',i,'Campos de '+name,[2.1,2.3,2.1])
d.core_properties.title='Sprint 3: Pedidos, ventas y procesos avanzados'
d.core_properties.author='David García; Yordy Carreto; Daylin Loarca'
d.core_properties.subject='Análisis de Sistemas II — Librería La Crayola Gigante'
out=ROOT/'docs/Sprint_3_Analisis_de_Sistemas.docx';d.save(out)
print(str(out));print('Tablas:',nt,'Figuras:',nf,'Párrafos:',len(d.paragraphs))
